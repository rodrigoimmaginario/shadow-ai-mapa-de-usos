"""Generate a blank public template and separated synthetic examples. No network."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
FIELDS = [
    ('ID do uso', 'Identificador local, sem identificação pessoal desnecessária.'),
    ('Área', 'Área responsável pelo trabalho; delimite o alcance do levantamento.'),
    ('Tarefa e benefício esperado', 'Descreva a tarefa. Diferencie expectativa de benefício medido.'),
    ('Ferramenta', 'Nome e modalidade. Não inclua credenciais ou URLs privadas.'),
    ('Tipo de conta', 'Conta corporativa não equivale a qualquer uso autorizado.'),
    ('Categorias de dados', 'Categorias apenas. Não cole dados reais, logs, código ou segredos.'),
    ('Capacidade / acesso', 'O que a ferramenta pode acessar ou fazer? Não é nível de maturidade.'),
    ('Situação de autorização', 'Informação declarada sobre este uso; não verificada automaticamente.'),
    ('Verificação do resultado', 'Quem verifica, como e quando? Diferencie ausência de desconhecimento.'),
    ('Lacunas a confirmar', 'Condições que precisam de verificação. Não cole conteúdo sensível.'),
    ('Próximo passo', 'Ação concreta. Suspeita de incidente segue o processo próprio da organização.'),
    ('Responsável pelo próximo passo', 'Papel ou referência interna responsável pelo encaminhamento.'),
    ('Data de revisão', 'Data acordada para revisar o uso. Não é prazo de resposta a incidente.'),
]
OPTIONS = {
    'E': ['Corporativa', 'Pessoal', 'Mista', 'Não se aplica', 'Não informado'],
    'G': ['Sem integração: recebe material e responde', 'Leitura de recurso conectado',
          'Propõe alterações sem aplicar', 'Altera ou executa ações', 'A confirmar'],
    'H': ['Autorizado para este uso', 'Ainda não avaliado', 'Não autorizado', 'A confirmar'],
}
EXAMPLES = [
    ['FIC-001', 'Desenvolvimento (fictício)',
     'Analisar log de integração; espera-se reduzir tempo de diagnóstico, ainda não medido.',
     'Assistente de IA genérico (fictício)', 'Pessoal',
     'Identificadores de cliente e informações técnicas — categorias fictícias, sem conteúdo real.',
     OPTIONS['G'][0], 'A confirmar',
     'Revisão técnica e testes antes da correção; condição fictícia a confirmar na empresa real.',
     'Regras para os dados e condições do serviço não confirmadas.',
     'Confirmar condições e alternativa aprovada; suspeita de exposição segue processo de segurança.',
     'Liderança técnica com segurança (papéis fictícios)', None],
    ['FIC-002', 'Comunicação (fictício)',
     'Rascunhar descrição pública de funcionalidade; benefício ainda não medido.',
     'Editor com IA (fictício)', 'Corporativa', 'Material já público, sem dados de clientes.',
     OPTIONS['G'][0], 'Autorizado para este uso',
     'Responsável pelo conteúdo revisa fatos e linguagem antes de publicar.',
     'Benefício e retrabalho ainda não acompanhados; autorização declarada no exemplo.',
     'Observar qualidade e esforço de revisão em um conjunto delimitado de entregas.',
     'Responsável pelo conteúdo (papel fictício)', None],
    ['FIC-003', 'Engenharia (fictício)',
     'Preparar correção em ambiente de testes; avaliar qualidade antes de ampliar acesso.',
     'Assistente conectado ao repositório (fictício)', 'Não informado',
     'Código interno; confirmar se configurações ou segredos podem ser acessados.',
     OPTIONS['G'][3], 'Ainda não avaliado',
     'Fluxo de revisão e aplicação de mudanças ainda a confirmar.',
     'Conta, permissões efetivas, ambiente alcançado e condições de parada.',
     'Não ampliar permissões; encaminhar avaliação de acessos e delimitar teste aprovado.',
     'Engenharia com segurança (papéis fictícios)', None],
]
NAVY = '173B58'
TEAL = '147D92'


def heading_row(ws, row, labels):
    for col, label in enumerate(labels, 1):
        cell = ws.cell(row, col, label)
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[row].height = 43


def style_inventory(ws, header_row, end_row):
    heading_row(ws, header_row, [f[0] for f in FIELDS])
    widths = [15, 23, 48, 32, 23, 48, 42, 30, 48, 48, 48, 34, 19]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(header_row, col).column_letter].width = width
        ws.cell(header_row, col).comment = Comment(FIELDS[col-1][1], 'Guia do modelo')
    for row in ws.iter_rows(min_row=header_row+1, max_row=end_row, max_col=13):
        for cell in row:
            cell.font = Font(name='Calibri', size=11)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        row[-1].number_format = 'dd/mm/yyyy'
        ws.row_dimensions[row[0].row].height = 75
    ws.sheet_view.zoomScale = 85
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = f'{header_row}:{header_row}'
    ws.oddFooter.center.text = 'Modelo v1.0 — preenchimento interno e confidencial'


def add_table(ws, name, ref):
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(table)


def build_workbook(dest: Path):
    wb = Workbook()
    start = wb.active
    start.title = 'Comece aqui'
    start.sheet_view.showGridLines = False
    start.column_dimensions['A'].width = 37
    start.column_dimensions['B'].width = 100
    rows = [
        ('MAPA INICIAL DE USOS DE IA', 'Versão 1.0 — Rodrigo Immaginario'),
        ('Finalidade', 'Conhecer usos declarados, preservar dúvidas e definir encaminhamentos. Sem nota de maturidade.'),
        ('Unidade / setor', None), ('Responsável pelo levantamento', None),
        ('Data do levantamento', None), ('Cobertura / o que ficou fora', None),
        ('Como começar', 'Ouça quem conhece o trabalho. Uma linha por tarefa + ferramenta + condições de dados e acesso.'),
        ('Onde preencher', 'Use Mapa de usos. Exemplo fictício é apenas demonstração e não integra seus registros.'),
        ('Como registrar dúvidas', 'Não informado / A confirmar não significam que algo não existe. Não preencha por suposição.'),
        ('O que não inserir', 'Não cole logs, código, dados pessoais, nomes de clientes, contratos, chaves, senhas ou tokens.'),
        ('Encaminhamento', 'Combine próximo passo, responsável e data de revisão. O modelo não define prioridade automaticamente.'),
        ('Incidentes', 'Suspeitas de exposição devem seguir o processo de segurança da empresa, sem esperar a revisão da planilha.'),
        ('Autorização', 'Registrar não autoriza. Conta corporativa não comprova autorização para qualquer uso.'),
        ('Privacidade', 'Planilha preenchida é interna. Não enviar ao autor nem publicar no GitHub, em comentários ou issues.'),
        ('Armazenamento', 'Use local autorizado pela organização. Controle acesso e retenção; o arquivo pode revelar lacunas de segurança.'),
        ('Limites', 'Não descobre ferramentas automaticamente, não audita contratos e não certifica conformidade ou segurança.'),
        ('Cobertura', 'Planilha vazia não comprova ausência de Shadow AI. Não generalize um setor para a empresa inteira.'),
        ('Mais linhas', 'O modelo oferece 80 linhas. Expanda a tabela e copie formatação/listas de uma linha vazia quando necessário.'),
        ('Listas de apoio', 'As opções são orientações de preenchimento, não mecanismos de bloqueio ou avaliação.'),
        ('Revisão', 'Revisite quando houver novos dados, mudanças de ferramenta, permissões ou processo.'),
        ('Documentação', 'Consulte a aba Guia de campos e https://rodrigoimmaginario.github.io/shadow-ai-mapa-de-usos/'),
        ('Distribuição', 'Modelo público para download e preenchimento interno. Não envie o arquivo preenchido ao autor ou ao GitHub.'),
    ]
    for row in rows:
        start.append(row)
    heading_row(start, 1, rows[0])
    for row in start.iter_rows(min_row=2):
        start.row_dimensions[row[0].row].height = 48
        for cell in row:
            cell.font = Font(name='Calibri', size=11, bold=cell.column == 1)
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        if 3 <= row[0].row <= 6:
            row[1].fill = PatternFill('solid', fgColor='E9F5F7')
    start['B5'].number_format = 'dd/mm/yyyy'
    start.freeze_panes = 'B3'
    start.sheet_properties.pageSetUpPr.fitToPage = True
    start.page_setup.fitToWidth = 1
    start.page_setup.fitToHeight = 0
    start.print_area = 'A1:B22'

    inventory = wb.create_sheet('Mapa de usos')
    style_inventory(inventory, 1, 81)
    inventory.freeze_panes = 'D2'
    add_table(inventory, 'MapaUsos', 'A1:M81')
    for column, choices in OPTIONS.items():
        dv = DataValidation(type='list', formula1='"' + ','.join(choices) + '"', allow_blank=True)
        dv.showErrorMessage = False
        dv.showInputMessage = True
        dv.promptTitle = 'Escolha ou esclareça a dúvida'
        dv.prompt = 'Não invente uma condição. A lista é apoio, não validação do uso.'
        dv.showDropDown = False
        inventory.add_data_validation(dv)
        dv.add(f'{column}2:{column}81')
    date = DataValidation(type='date', operator='between', formula1='36526', formula2='73415', allow_blank=True)
    date.showInputMessage = True
    date.promptTitle = 'Revisão acordada'
    date.prompt = 'Data combinada com o responsável. Não é prazo de resposta a incidente.'
    date.showErrorMessage = True
    date.errorTitle = 'Use uma data válida'
    date.error = 'Digite uma data, por exemplo 15/10/2026, ou deixe em branco até combinar.'
    inventory.add_data_validation(date)
    date.add('M2:M81')

    example = wb.create_sheet('Exemplo fictício')
    example.merge_cells('A1:M1')
    example['A1'] = 'EXEMPLOS FICTÍCIOS — não são registros da organização, incidentes reais ou resultados medidos'
    example['A1'].font = Font(name='Calibri', bold=True, color=NAVY, size=13)
    example.row_dimensions[1].height = 32
    style_inventory(example, 3, 6)
    for row, values in enumerate(EXAMPLES, 4):
        for col, value in enumerate(values, 1):
            example.cell(row, col, value)
        example.row_dimensions[row].height = 112
    add_table(example, 'ExemplosFicticios', 'A3:M6')
    example.freeze_panes = 'D4'

    guide = wb.create_sheet('Guia de campos')
    heading_row(guide, 1, ['Campo', 'Orientação'])
    guide.column_dimensions['A'].width = 38
    guide.column_dimensions['B'].width = 110
    for field in FIELDS:
        guide.append(field)
    for row in guide.iter_rows(min_row=2):
        guide.row_dimensions[row[0].row].height = 42
        for cell in row:
            cell.font = Font(name='Calibri', size=11)
            cell.alignment = Alignment(wrap_text=True, vertical='center')
    guide.freeze_panes = 'A2'
    guide.sheet_view.showGridLines = False
    wb.properties.creator = 'Rodrigo Immaginario'
    wb.properties.lastModifiedBy = 'Rodrigo Immaginario'
    wb.properties.title = 'Mapa inicial de usos de IA — v1.0'
    wb.properties.description = 'Levantamento declarado; sem notas de maturidade. Exemplos fictícios separados.'
    wb.save(dest)
    wb.close()


if __name__ == "__main__":
    target = ROOT / "docs/downloads/mapa-usos-ia-v1.xlsx"
    target.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(target)
    print(target)
