import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile
from bs4 import BeautifulSoup
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parent.parent
XLSX=ROOT/'docs/downloads/mapa-usos-ia-v1.xlsx'

class PublicKitTests(unittest.TestCase):
    def setUp(self):self.wb=load_workbook(XLSX)
    def tearDown(self):self.wb.close()
    def test_tabs_and_empty_inventory(self):
        self.assertEqual(self.wb.sheetnames,['Comece aqui','Mapa de usos','Exemplo fictício','Guia de campos'])
        ws=self.wb['Mapa de usos']
        self.assertEqual(ws.tables['MapaUsos'].ref,'A1:M81')
        self.assertEqual(ws.freeze_panes,'D2')
        self.assertTrue(all(c.value is None for row in ws.iter_rows(min_row=2,max_row=81,max_col=13) for c in row))
    def test_no_formulas_connections_or_macros(self):
        self.assertFalse(any(c.data_type=='f' for ws in self.wb for row in ws for c in row))
        with ZipFile(XLSX) as z:
            self.assertIsNone(z.testzip())
            self.assertFalse(any('vbaProject' in n or 'externalLinks/' in n or 'connections.xml' in n for n in z.namelist()))
    def test_separate_fictional_examples(self):
        ws=self.wb['Exemplo fictício']
        self.assertIn('FICTÍCIOS',ws['A1'].value)
        self.assertEqual([ws.cell(i,1).value for i in range(4,7)],['FIC-001','FIC-002','FIC-003'])
        self.assertTrue(all(ws.cell(i,13).value is None for i in range(4,7)))
        self.assertEqual(ws['H4'].value,'A confirmar')
    def test_dropdowns_and_private_use_warning(self):
        self.assertEqual(len(self.wb['Mapa de usos'].data_validations.dataValidation),4)
        texts=' '.join(str(c.value or '') for row in self.wb['Comece aqui'] for c in row)
        self.assertIn('Não enviar ao autor',texts)
        self.assertIn('Sem nota de maturidade',texts)
    def test_author_metadata(self):
        self.assertEqual(self.wb.properties.creator,'Rodrigo Immaginario')
        self.assertNotIn('pré-produção',self.wb.properties.title.lower())
    def test_fill_save_reopen(self):
        ws=self.wb['Mapa de usos'];ws['A2']='FIC-TESTE';ws['E2']='Pessoal';ws['H2']='A confirmar'
        with tempfile.TemporaryDirectory() as temp:
            p=Path(temp)/'test.xlsx';self.wb.save(p)
            other=load_workbook(p)
            try:self.assertEqual(other['Mapa de usos']['H2'].value,'A confirmar')
            finally:other.close()
    def test_static_page_and_download(self):
        soup=BeautifulSoup((ROOT/'docs/index.html').read_text(encoding='utf-8'),'html.parser')
        self.assertEqual(soup.html['lang'],'pt-BR')
        self.assertFalse(soup.find_all(['script','form','iframe']))
        self.assertFalse(soup.select('[src],link[rel="stylesheet"]'))
        downloads=[a['href'] for a in soup.find_all('a',href=True) if a['href'].endswith('.xlsx')]
        self.assertEqual(downloads,['downloads/mapa-usos-ia-v1.xlsx'])
        self.assertTrue((ROOT/'docs'/downloads[0]).is_file())
        self.assertIn('Não envie respostas',soup.get_text())
        self.assertNotIn('C:/',soup.get_text())
        self.assertNotIn('ainda não publicado',soup.get_text())

if __name__=='__main__':unittest.main()
